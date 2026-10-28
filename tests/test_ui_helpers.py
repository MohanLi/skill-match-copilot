from __future__ import annotations

import csv
import io

import pytest

from src.ui.helpers import highlight, read_resume, results_csv, safe_url, search_key


def test_text_resume_preserves_unicode_and_removes_bom():
    assert read_resume('\ufeffPython, SQL\n研究'.encode(), 'resume.txt') == 'Python, SQL\n研究'


def test_docx_resume_keeps_paragraph_and_table_order():
    from docx import Document

    doc = Document()
    doc.add_paragraph('Experience')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Python'
    table.cell(0, 1).text = 'Docker'
    doc.add_paragraph('Education')
    data = io.BytesIO()
    doc.save(data)
    assert read_resume(data.getvalue(), 'resume.docx') == 'Experience\nPython | Docker\nEducation'


def test_pdf_resume_reads_selectable_text():
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(612, 792)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'),
                             NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 30 700 Td (Python and Docker) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    data = io.BytesIO()
    writer.write(data)
    assert read_resume(data.getvalue(), 'resume.pdf') == 'Python and Docker'


@pytest.mark.parametrize('data,name,message', [
    (b'', 'empty.txt', 'No text'),
    (b'hello', 'resume.exe', 'PDF, DOCX'),
    (b'\xff', 'resume.txt', 'UTF-8'),
    (b'x' * 20_001, 'resume.txt', '20,000'),
    (b'x' * (10 * 1024 * 1024 + 1), 'resume.pdf', '10 MB'),
    (b'not a docx', 'resume.docx', "couldn't read"),
])
def test_invalid_resume_has_actionable_error(data, name, message):
    with pytest.raises(ValueError, match=message):
        read_resume(data, name)


def test_highlight_escapes_text_and_does_not_nest_overlapping_quotes():
    text = 'Python & Docker <script>alert(1)</script>\npython'
    rendered = highlight(text, ['Python & Docker', 'Docker', 'Python', 'script'])
    assert '<script>' not in rendered
    assert 'Python &amp; Docker</mark>' in rendered
    assert rendered.count('<mark ') == rendered.count('</mark>') == 4
    assert '<br>' in rendered
    assert '<mark' not in highlight('<mark>', [])


@pytest.mark.parametrize('url,expected', [
    ('https://example.com/job', 'https://example.com/job'), (' http://example.com/job ', 'http://example.com/job'),
    ('javascript:alert(1)', None), ('data:text/html,test', None), ('https://[bad', None), ('/jobs/1', None), (None, None),
])
def test_job_links_require_http_or_https(url, expected):
    assert safe_url(url) == expected


def test_search_key_changes_with_resume_and_settings():
    values = ['resume', 'remote', 'live_jobs', 'single', 5, 'agent']
    original = search_key(*values)
    for i, replacement in enumerate(['edited resume', 'Chicago', 'djinni_jobs', 'baseline', 10, 'another-model']):
        edited = values.copy()
        edited[i] = replacement
        assert search_key(*edited) != original


def test_csv_scores_and_formula_like_titles_are_labelled_safely():
    result = {'mode': 'baseline', 'run': {'ranking': [{'job_id': '1', 'rank': 1, 'fit_score': 75}]},
              'jobs': {'1': {'title': '=HYPERLINK("bad")', 'company': 'ACME', 'url': 'javascript:alert(1)'}}}
    row = list(csv.DictReader(io.StringIO(results_csv(result))))[0]
    assert row['title'].startswith("'=")
    assert row['score_type'] == 'cosine_similarity_x100'
    assert row['url'] == ''
